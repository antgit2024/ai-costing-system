from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ...database import get_db
from ..models import TmallSkuGeneratorTemplate

router = APIRouter(prefix="/tmall/sku-template", tags=["tmall"])


# -----------------------------
# Schemas (MVP)
# -----------------------------


class TmallSizeOption(BaseModel):
    key: str = Field(..., description="尺寸key（内部唯一，例如 size_1）")
    label: str = Field(..., description="尺寸展示值（例如：枕套 / 枕芯+枕套）")
    # Optional per-size defaults
    size_code: Optional[str] = Field(None, description="用于商家编码的尺寸段（例如 C1/C2）")


class TmallColorOption(BaseModel):
    key: str = Field(..., description="颜色分类key（内部唯一）")
    label: str = Field(..., description="颜色分类展示值（天猫语义=图案/工艺/款式）")
    width_cm: Optional[float] = Field(None, description="宽(cm)")
    height_cm: Optional[float] = Field(None, description="高(cm)")
    thickness_cm: Optional[float] = Field(None, description="厚度(cm)")
    length_cm: Optional[float] = Field(None, description="长度(cm)")
    main_pattern_type: Optional[str] = Field(None, description="主图案类型（可选）")


class TmallSkuCell(BaseModel):
    color_key: str
    size_key: str
    enabled: bool = True
    # Optional override
    merchant_sku: Optional[str] = None


class TmallSkuTemplateRequest(BaseModel):
    # Basic
    sizes: List[TmallSizeOption]
    colors: List[TmallColorOption]
    cells: List[TmallSkuCell]
    # Defaults
    default_enabled: bool = True
    default_sku_status: int = 1

    # Merchant SKU generator (simple MVP)
    merchant_sku_prefix: str = Field("", description="商家编码前缀，例如 BZPB008XXXXX-")
    merchant_sku_sep: str = Field("-", description="分隔符")
    merchant_sku_suffix: str = Field("", description="商家编码后缀（可选）")


class TmallSkuRow(BaseModel):
    color_label: str
    size_label: str
    merchant_sku: str
    sku_status: int
    main_pattern_type: Optional[str] = None
    length_cm: Optional[str] = None
    thickness_cm: Optional[str] = None
    width_cm: Optional[str] = None


class TmallSkuTemplatePreviewResponse(BaseModel):
    total_rows: int
    rows: List[TmallSkuRow]
    # Sheet2 mapping (for transparency)
    header_mapping: Dict[str, str]


class TmallSkuGeneratorPersistedConfig(BaseModel):
    merchantSkuPrefix: str = ""
    merchantSkuSuffix: str = ""
    listingChannel: Optional[str] = None
    sizes: List[Dict[str, Any]] = Field(default_factory=list)
    colors: List[Dict[str, Any]] = Field(default_factory=list)
    mainPatternTypes: List[Dict[str, Any]] = Field(default_factory=list)
    ui: Dict[str, Any] = Field(default_factory=dict)

    class Config:
        extra = "allow"


class TmallSkuGeneratorTemplateRead(BaseModel):
    id: str
    name: str
    type: str
    published_at: Optional[datetime] = None
    matrix_count: Optional[int] = None
    archived: bool = False
    config: TmallSkuGeneratorPersistedConfig
    created_at: datetime
    updated_at: datetime


class TmallSkuGeneratorTemplateListResponse(BaseModel):
    total: int
    items: List[TmallSkuGeneratorTemplateRead]


class TmallSkuGeneratorTemplateUpsertRequest(BaseModel):
    id: str = Field(..., min_length=1, max_length=64)
    name: str = Field("未命名模板", max_length=255)
    type: str = Field("家居布艺", max_length=64)
    published_at: Optional[datetime] = None
    matrix_count: Optional[int] = None
    archived: bool = False
    config: TmallSkuGeneratorPersistedConfig = Field(default_factory=TmallSkuGeneratorPersistedConfig)


class TmallSkuGeneratorTemplateUpdateRequest(BaseModel):
    name: Optional[str] = Field(None, max_length=255)
    type: Optional[str] = Field(None, max_length=64)
    published_at: Optional[datetime] = None
    matrix_count: Optional[int] = None
    archived: Optional[bool] = None
    config: Optional[TmallSkuGeneratorPersistedConfig] = None


TMALL_SHEET2_MAPPING: Dict[str, str] = {
    "颜色分类": "p-1627207",
    "尺寸": "p-21433",
    "商家编码": "skuOuterId",
    "是否上架": "skuStatus",
    "主图案类型": "skuParam_p-20603",
    "长度": "skuParam_p-554361099",
    "厚度(cm)": "skuParam_p-554360987",
    "宽度": "skuParam_p-554668632",
}


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _default_generator_config() -> Dict[str, Any]:
    return {
        "merchantSkuPrefix": "BZPB008XXXXX-",
        "merchantSkuSuffix": "",
        "listingChannel": "tmall",
        "sizes": [
            {"key": "size_1", "label": "枕芯+枕套", "size_code": "C1", "source_code": ""},
            {"key": "size_2", "label": "枕套", "size_code": "C2", "source_code": ""},
        ],
        "colors": [
            {
                "key": "c1",
                "label": "Q25122501A黄金绒背面纯色（红色毛球） 45X45",
                "width_cm": 45,
                "height_cm": 45,
                "enabledSizes": {"size_1": True, "size_2": True},
                "source_code": "",
            }
        ],
        "mainPatternTypes": [{"key": "p1", "label": "无"}],
        "ui": {
            "enableColorImages": True,
            "enableSizeImages": False,
            "enableColorRemarks": True,
            "enableSizeRemarks": True,
            "enablePatternRemarks": False,
            "includeMainPatternType": True,
        },
    }


def _serialize_generator_template(row: TmallSkuGeneratorTemplate) -> TmallSkuGeneratorTemplateRead:
    return TmallSkuGeneratorTemplateRead(
        id=row.id,
        name=row.name,
        type=row.type,
        published_at=row.published_at,
        matrix_count=row.matrix_count,
        archived=bool(row.is_archived),
        config=TmallSkuGeneratorPersistedConfig.parse_obj(row.config_json or {}),
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def _ensure_generator_seed(db: Session) -> None:
    exists = db.get(TmallSkuGeneratorTemplate, "mvp")
    if exists:
        return
    now = _utcnow()
    db.add(
        TmallSkuGeneratorTemplate(
            id="mvp",
            name="天猫布艺 SKU规格生成器（MVP）",
            type="家居布艺",
            published_at=now,
            matrix_count=2,
            config_json=_default_generator_config(),
        )
    )
    db.commit()


def _get_generator_template_or_404(db: Session, template_id: str) -> TmallSkuGeneratorTemplate:
    row = db.get(TmallSkuGeneratorTemplate, template_id)
    if not row:
        raise HTTPException(status_code=404, detail="Tmall SKU generator template not found")
    return row


@router.get("/generator-templates", response_model=TmallSkuGeneratorTemplateListResponse)
def list_generator_templates(
    search: Optional[str] = None,
    type: Optional[str] = None,
    include_archived: bool = False,
    db: Session = Depends(get_db),
) -> TmallSkuGeneratorTemplateListResponse:
    _ensure_generator_seed(db)
    q = db.query(TmallSkuGeneratorTemplate)
    if not include_archived:
        q = q.filter(TmallSkuGeneratorTemplate.is_archived.is_(False))
    if type:
        q = q.filter(TmallSkuGeneratorTemplate.type == type)
    if search:
        like = f"%{search.strip()}%"
        q = q.filter(or_(TmallSkuGeneratorTemplate.id.ilike(like), TmallSkuGeneratorTemplate.name.ilike(like)))
    items = q.order_by(TmallSkuGeneratorTemplate.updated_at.desc()).all()
    return TmallSkuGeneratorTemplateListResponse(total=len(items), items=[_serialize_generator_template(x) for x in items])


@router.post("/generator-templates", response_model=TmallSkuGeneratorTemplateRead, status_code=status.HTTP_201_CREATED)
def create_generator_template(
    payload: TmallSkuGeneratorTemplateUpsertRequest,
    db: Session = Depends(get_db),
) -> TmallSkuGeneratorTemplateRead:
    template_id = payload.id.strip()
    if db.get(TmallSkuGeneratorTemplate, template_id):
        raise HTTPException(status_code=409, detail="Template id already exists")
    now = payload.published_at or _utcnow()
    row = TmallSkuGeneratorTemplate(
        id=template_id,
        name=(payload.name or "未命名模板").strip() or "未命名模板",
        type=(payload.type or "家居布艺").strip() or "家居布艺",
        published_at=now,
        matrix_count=payload.matrix_count,
        is_archived=payload.archived,
        config_json=payload.config.dict(),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return _serialize_generator_template(row)


@router.get("/generator-templates/{template_id}", response_model=TmallSkuGeneratorTemplateRead)
def get_generator_template(template_id: str, db: Session = Depends(get_db)) -> TmallSkuGeneratorTemplateRead:
    _ensure_generator_seed(db)
    return _serialize_generator_template(_get_generator_template_or_404(db, template_id))


@router.put("/generator-templates/{template_id}", response_model=TmallSkuGeneratorTemplateRead)
def upsert_generator_template(
    template_id: str,
    payload: TmallSkuGeneratorTemplateUpsertRequest,
    db: Session = Depends(get_db),
) -> TmallSkuGeneratorTemplateRead:
    tid = template_id.strip()
    row = db.get(TmallSkuGeneratorTemplate, tid)
    if not row:
        row = TmallSkuGeneratorTemplate(id=tid)
        db.add(row)
    row.name = (payload.name or "未命名模板").strip() or "未命名模板"
    row.type = (payload.type or "家居布艺").strip() or "家居布艺"
    row.published_at = payload.published_at or _utcnow()
    row.matrix_count = payload.matrix_count
    row.is_archived = payload.archived
    row.config_json = payload.config.dict()
    db.commit()
    db.refresh(row)
    return _serialize_generator_template(row)


@router.patch("/generator-templates/{template_id}", response_model=TmallSkuGeneratorTemplateRead)
def update_generator_template(
    template_id: str,
    payload: TmallSkuGeneratorTemplateUpdateRequest,
    db: Session = Depends(get_db),
) -> TmallSkuGeneratorTemplateRead:
    row = _get_generator_template_or_404(db, template_id)
    if payload.name is not None:
        row.name = payload.name.strip() or "未命名模板"
    if payload.type is not None:
        row.type = payload.type.strip() or "家居布艺"
    if payload.published_at is not None:
        row.published_at = payload.published_at
    if payload.matrix_count is not None:
        row.matrix_count = payload.matrix_count
    if payload.archived is not None:
        row.is_archived = payload.archived
    if payload.config is not None:
        row.config_json = payload.config.dict()
    db.commit()
    db.refresh(row)
    return _serialize_generator_template(row)


def _fmt_num(v: Optional[float]) -> Optional[str]:
    if v is None:
        return None
    try:
        n = float(v)
    except Exception:
        return None
    if n != n:  # NaN
        return None
    # keep up to 2 decimals, strip trailing zeros
    s = f"{n:.2f}"
    s = s.rstrip("0").rstrip(".")
    return s


def _build_preview_rows(payload: TmallSkuTemplateRequest) -> List[TmallSkuRow]:
    size_by_key: Dict[str, TmallSizeOption] = {s.key: s for s in payload.sizes}
    color_by_key: Dict[str, TmallColorOption] = {c.key: c for c in payload.colors}

    rows: List[TmallSkuRow] = []

    for cell in payload.cells:
        c = color_by_key.get(cell.color_key)
        s = size_by_key.get(cell.size_key)
        if not c or not s:
            continue

        enabled = bool(cell.enabled)
        sku_status = 1 if enabled else 0

        # Merchant sku (MVP): prefix + WxH + size_code + suffix
        if cell.merchant_sku:
            merchant = str(cell.merchant_sku).strip()
        else:
            w = _fmt_num(c.width_cm) or ""
            h = _fmt_num(c.height_cm) or ""
            wh = ""
            if w and h:
                wh = f"{w}{h}"
            code = str(s.size_code or "").strip()
            merchant = f"{payload.merchant_sku_prefix}{wh}{code}{payload.merchant_sku_suffix}"
            merchant = merchant.strip() or ""

        rows.append(
            TmallSkuRow(
                color_label=c.label,
                size_label=s.label,
                merchant_sku=merchant,
                sku_status=sku_status,
                main_pattern_type=(c.main_pattern_type or None),
                length_cm=_fmt_num(c.length_cm),
                thickness_cm=_fmt_num(c.thickness_cm),
                width_cm=_fmt_num(c.width_cm),
            )
        )

    return rows


@router.post("/preview", response_model=TmallSkuTemplatePreviewResponse)
def preview_tmall_sku_template(payload: TmallSkuTemplateRequest) -> TmallSkuTemplatePreviewResponse:
    rows = _build_preview_rows(payload)
    return TmallSkuTemplatePreviewResponse(total_rows=len(rows), rows=rows, header_mapping=TMALL_SHEET2_MAPPING)


@router.post("/export")
def export_tmall_sku_template(payload: TmallSkuTemplateRequest) -> StreamingResponse:
    from openpyxl import Workbook

    rows = _build_preview_rows(payload)

    wb = Workbook(write_only=True)

    # Sheet1: data
    ws = wb.create_sheet("sheet1")
    headers = ["颜色分类", "尺寸", "商家编码", "是否上架", "主图案类型", "长度", "厚度(cm)", "宽度"]
    ws.append(headers)
    for r in rows:
        ws.append(
            [
                r.color_label,
                r.size_label,
                r.merchant_sku,
                str(r.sku_status),
                r.main_pattern_type or "",
                r.length_cm or "",
                r.thickness_cm or "",
                r.width_cm or "",
            ]
        )

    # Sheet2: mapping
    ws2 = wb.create_sheet("sheet2")
    ws2.append(["字段名", "平台字段key"])
    for k, v in TMALL_SHEET2_MAPPING.items():
        ws2.append([k, v])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    filename = "tmall_buyi_sku_template.xlsx"
    headers: Dict[str, Any] = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )

